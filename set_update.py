#!/usr/bin/env python3
"""Validate the precomputed SET signal workbook without placing orders."""

from __future__ import annotations

import math
import os
import re
import sys
import tempfile
import threading
import time
from datetime import datetime, timedelta, timezone
from typing import Any, Callable, Optional
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlparse
from urllib.request import Request, urlopen
from zipfile import BadZipFile

import openpyxl


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ENV_FILE = os.path.join(SCRIPT_DIR, ".env")
DOWNLOAD_DIR = os.path.join(SCRIPT_DIR, "downloads")
WORKBOOK_PATH = os.path.join(DOWNLOAD_DIR, "All-SET-MTFS.xlsx")
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
SETTRADE_REALTIME_TIMEOUT_SECONDS = 20
SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS = 120
SETTRADE_HISTORY_TRADING_DAYS = 20
SETTRADE_HISTORY_CANDLE_LIMIT = 40
THAILAND_TIMEZONE = timezone(timedelta(hours=7))

REQUIRED_COLUMNS = (
    "AKA",
    "GROUP",
    "PRICE",
    "TREND(W)",
    "RSI(W)",
    "MACD_HIST_DELTA(W)",
    "RELATIVE_STRENGTH(W)",
    "TREND(D)",
    "RSI(D)",
    "RELATIVE_STRENGTH(D)",
    "TREND(4H)",
    "RSI(4H)",
    "RELATIVE_STRENGTH(4H)",
    "TREND(1H)",
    "RSI(1H)",
    "RELATIVE_STRENGTH(1H)",
    "ACTION",
)
RSI_COLUMNS = ("RSI(W)", "RSI(D)", "RSI(4H)", "RSI(1H)")
TREND_COLUMNS = ("TREND(W)", "TREND(D)", "TREND(4H)", "TREND(1H)")
RELATIVE_STRENGTH_COLUMNS = (
    "RELATIVE_STRENGTH(W)",
    "RELATIVE_STRENGTH(D)",
    "RELATIVE_STRENGTH(4H)",
    "RELATIVE_STRENGTH(1H)",
)
SYMBOL_PATTERN = re.compile(r"[A-Z0-9][A-Z0-9.&_-]{0,19}\Z")
SHEET_ID_PATTERN = re.compile(r"/spreadsheets/d/([A-Za-z0-9_-]+)(?:/|$)")
TIMESTAMP_COLUMNS = {"UPDATED", "LAST_CANDLE", "AS_OF", "TIMESTAMP", "DATE"}
UPDATE_SUMMARY_SYMBOL = "ADVANC"
CANDIDATE_COLUMNS = (
    "TREND(W)",
    "TREND(D)",
    "RSI(W)",
    "MACD_HIST_DELTA(W)",
    "RSI(D)",
    "EMA9(W)",
    "EMA21(W)",
    "SIGNAL_VOLUME(W)",
    "SIGNAL_VOLUME(D)",
    "SIGNAL_VOLUME(4H)",
    "SIGNAL_VOLUME(1H)",
)
REPORT_COLUMNS = CANDIDATE_COLUMNS + ("ENTRY_ZONE(4H)", "ENTRY_TRIGGER(1H)")


def finnomena_stock_url(symbol: str) -> str:
    return f"https://www.finnomena.com/stock/{quote(symbol, safe='')}/advance"


def load_config() -> dict[str, str]:
    config = {
        key: value
        for key, value in os.environ.items()
        if key in {
            "SET_DATA_URL",
            "SET_DATA_SHEET",
            "SETTRADE_APP_ID",
            "SETTRADE_APP_SECRET",
            "SETTRADE_BROKER_ID",
            "SETTRADE_APP_CODE",
        }
    }
    try:
        with open(ENV_FILE, encoding="utf-8") as env_file:
            for raw_line in env_file:
                line = raw_line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                key, value = line.split("=", 1)
                key = key.strip()
                value = value.strip()
                if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
                    value = value[1:-1]
                if key in {
                    "SET_DATA_URL",
                    "SET_DATA_SHEET",
                    "SETTRADE_APP_ID",
                    "SETTRADE_APP_SECRET",
                    "SETTRADE_BROKER_ID",
                    "SETTRADE_APP_CODE",
                }:
                    config.setdefault(key, value)
    except FileNotFoundError as exc:
        if not config.get("SET_DATA_URL", "").strip():
            raise ValueError("ไม่พบไฟล์ .env; กรุณาคัดลอก .env.example เป็น .env และกำหนด SET_DATA_URL") from exc
    except OSError as exc:
        raise ValueError(f"อ่านไฟล์ .env ไม่สำเร็จ: {exc}") from exc

    config.setdefault("SET_DATA_SHEET", "All-SET-MTFS")
    if not config.get("SET_DATA_URL", "").strip():
        raise ValueError("กรุณากำหนด SET_DATA_URL ในไฟล์ .env")
    return config


def spreadsheet_export_url(sheet_url: str) -> str:
    parsed = urlparse(sheet_url)
    match = SHEET_ID_PATTERN.search(parsed.path)
    if parsed.scheme != "https" or parsed.hostname != "docs.google.com" or not match:
        raise ValueError("SET_DATA_URL ต้องเป็น URL ของ Google Sheets ที่เข้าถึงได้ด้วยลิงก์")
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx"


def download_workbook(url: str) -> bytes:
    request = Request(
        url,
        headers={
            "User-Agent": "TFOSvS/SET-data-validator",
            "Cache-Control": "no-cache, no-store, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )
    try:
        with urlopen(request, timeout=30) as response:
            content = response.read(MAX_WORKBOOK_BYTES + 1)
    except HTTPError as exc:
        raise ValueError(f"ดาวน์โหลด Google Sheets ไม่สำเร็จ (HTTP {exc.code})") from exc
    except URLError as exc:
        raise ValueError(f"เชื่อมต่อ Google Sheets ไม่สำเร็จ: {exc.reason}") from exc
    except TimeoutError as exc:
        raise ValueError("หมดเวลาขณะดาวน์โหลด Google Sheets") from exc

    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError(f"ไฟล์ workbook มีขนาดเกิน {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB")

    os.makedirs(DOWNLOAD_DIR, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", dir=DOWNLOAD_DIR, prefix=".All-SET-MTFS-", suffix=".tmp", delete=False
        ) as temporary_file:
            temporary_path = temporary_file.name
            temporary_file.write(content)
        os.replace(temporary_path, WORKBOOK_PATH)
    except OSError as exc:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)
        raise ValueError(f"บันทึก workbook ลง downloads ไม่สำเร็จ: {exc}") from exc
    return content


def validate_worksheet(worksheet) -> tuple[int, list[str], bool, int]:
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        if header and header not in indexes:
            indexes[header] = index

    missing_columns = [column for column in REQUIRED_COLUMNS if column not in indexes]
    if missing_columns:
        message = f"ขาดคอลัมน์ที่จำเป็น: {', '.join(missing_columns)}"
        return 0, [message], False, 1

    timestamp_available = {"LAST_CANDLE", "UPDATED"}.issubset(indexes)
    errors: list[str] = []
    error_count = 0
    symbols: set[str] = set()
    row_count = 0

    def add_error(message: str) -> None:
        nonlocal error_count
        error_count += 1
        if len(errors) < 100:
            errors.append(message)

    for excel_row, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        row_count += 1
        raw_symbol = row[indexes["AKA"]] if indexes["AKA"] < len(row) else None
        symbol = "TRUE" if raw_symbol is True else (
            raw_symbol.strip().upper() if isinstance(raw_symbol, str) else ""
        )
        if not SYMBOL_PATTERN.fullmatch(symbol):
            add_error(f"แถว {excel_row}: AKA ว่างหรือมีรูปแบบไม่ถูกต้อง")
        elif symbol in symbols:
            add_error(f"แถว {excel_row}: AKA ซ้ำ")
        else:
            symbols.add(symbol)

        group_index = indexes["GROUP"]
        group = row[group_index] if group_index < len(row) else None
        if group is None or not str(group).strip():
            add_error(f"แถว {excel_row}: GROUP ว่าง")

        for column in TREND_COLUMNS + RELATIVE_STRENGTH_COLUMNS:
            value = row[indexes[column]] if indexes[column] < len(row) else None
            if value is None or not str(value).strip():
                add_error(f"แถว {excel_row}: {column} ว่าง")

        for column in ("PRICE", "MACD_HIST_DELTA(W)") + RSI_COLUMNS:
            value = row[indexes[column]] if indexes[column] < len(row) else None
            if isinstance(value, bool):
                add_error(f"แถว {excel_row}: {column} ไม่ใช่ตัวเลข")
                continue
            try:
                number = float(value)
            except (TypeError, ValueError, OverflowError):
                add_error(f"แถว {excel_row}: {column} ไม่ใช่ตัวเลข")
                continue
            if not math.isfinite(number):
                add_error(f"แถว {excel_row}: {column} ไม่ใช่ค่าตัวเลข finite")
            elif column == "PRICE" and number <= 0:
                add_error(f"แถว {excel_row}: PRICE ต้องมากกว่า 0")
            elif column in RSI_COLUMNS and not 0 <= number <= 100:
                add_error(f"แถว {excel_row}: {column} ต้องอยู่ระหว่าง 0 ถึง 100")

    if row_count == 0:
        add_error("ไม่พบแถวข้อมูลในชีต")
    return row_count, errors, timestamp_available, error_count


def _header_indexes(worksheet) -> dict[str, int]:
    headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {
        str(value).strip(): index
        for index, value in enumerate(headers)
        if value is not None and str(value).strip()
    }


def _cell(row, indexes: dict[str, int], column: str):
    index = indexes[column]
    return row[index] if index < len(row) else None


def _display(value) -> str:
    return str(value).strip() if value is not None else ""


def _update_summary_row(worksheet, indexes: dict[str, int]):
    required_columns = ("AKA", "LAST_CANDLE", "UPDATED", "GROUP")
    missing = [column for column in required_columns if column not in indexes]
    if missing:
        raise ValueError("ขาดคอลัมน์สำหรับสรุป UPDATE: " + ", ".join(missing))

    target_row = next(
        (
            row for row in worksheet.iter_rows(min_row=2, values_only=True)
            if _display(_cell(row, indexes, "AKA")).upper() == UPDATE_SUMMARY_SYMBOL
        ),
        None,
    )
    if target_row is None:
        raise ValueError(f"ไม่พบ AKA {UPDATE_SUMMARY_SYMBOL} ใน snapshot")

    timestamp_values = {
        column: _display(_cell(target_row, indexes, column))
        for column in ("LAST_CANDLE", "UPDATED")
    }
    missing_values = [column for column, value in timestamp_values.items() if not value]
    if missing_values:
        raise ValueError(
            f"แถว {UPDATE_SUMMARY_SYMBOL}: ไม่มีค่า " + ", ".join(missing_values)
        )
    return target_row


def _print_update_summary(target_row, indexes: dict[str, int]) -> None:
    timestamp_values = {
        column: _display(_cell(target_row, indexes, column))
        for column in ("LAST_CANDLE", "UPDATED")
    }
    print("\nสรุปข้อมูลหุ้น")
    print(f"{'AKA':<12}: {UPDATE_SUMMARY_SYMBOL} 🔔")
    print(f"{'LAST_CANDLE':<12}: {timestamp_values['LAST_CANDLE']}")
    print(f"{'UPDATED':<12}: {timestamp_values['UPDATED']}")
    print(f"{'GROUP':<12}: {_display(_cell(target_row, indexes, 'GROUP'))}")


def _markdown_cell(value) -> str:
    return _display(value).replace("|", r"\|").replace("\n", " ")


def _format_price(value) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    return _display(value)


def _format_price_fixed(value) -> str:
    number = _finite_number(value)
    return f"{number:,.4f}" if number is not None else _display(value)


def _finite_number(value) -> Optional[float]:
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def _parse_market_message(message: dict[str, Any], kind: str) -> dict[str, Optional[float]]:
    if not message.get("is_success"):
        raise ValueError(str(message.get("message") or "Settrade ส่งข้อมูลไม่สำเร็จ"))
    data = message.get("data")
    if not isinstance(data, dict):
        raise ValueError("รูปแบบข้อมูลจาก Settrade ไม่ถูกต้อง")

    if kind == "info":
        value = _finite_number(data.get("total_value"))
        if value is None or value < 0:
            raise ValueError("ไม่พบมูลค่าซื้อขายสะสมวันนี้")
        return {"turnover": value}
    if kind != "book":
        raise ValueError(f"ไม่รองรับข้อมูลตลาดประเภท {kind}")

    bid_depth = 0.0
    ask_depth = 0.0
    for level in range(1, 11):
        bid_volume = _finite_number(data.get(f"bid_volume{level}"))
        ask_volume = _finite_number(data.get(f"ask_volume{level}"))
        if bid_volume is None or ask_volume is None or bid_volume < 0 or ask_volume < 0:
            raise ValueError("ข้อมูลปริมาณ bid/ask 10 ระดับไม่ครบหรือไม่ถูกต้อง")
        bid_depth += bid_volume
        ask_depth += ask_volume

    bid_price = _finite_number(data.get("bid_price1"))
    ask_price = _finite_number(data.get("ask_price1"))
    if (
        bid_price is None or ask_price is None
        or bid_price < 0 or ask_price < 0
    ):
        raise ValueError("ข้อมูลราคา bid/ask ระดับแรกไม่ครบหรือไม่ถูกต้อง")
    spread = (
        ask_price - bid_price
        if bid_price > 0 and ask_price > 0
        else None
    )
    if spread is not None and spread < 0:
        raise ValueError("ราคาเสนอซื้อ/ขายระดับแรกไม่ถูกต้อง")
    return {
        "bid_price": bid_price,
        "ask_price": ask_price,
        "spread": spread,
        "bid_depth": bid_depth,
        "ask_depth": ask_depth,
    }


def _collect_settrade_liquidity(
    symbols: list[str],
    config: dict[str, str],
    *,
    investor_factory: Optional[Callable[..., Any]] = None,
    max_duration_seconds: Optional[float] = None,
) -> tuple[dict[str, dict[str, Optional[float]]], list[str]]:
    if not symbols:
        return {}, []
    app_id = config.get("SETTRADE_APP_ID", "").strip()
    app_secret = config.get("SETTRADE_APP_SECRET", "").strip()
    if not app_id or not app_secret:
        missing = [
            name for name, value in (
                ("SETTRADE_APP_ID", app_id),
                ("SETTRADE_APP_SECRET", app_secret),
            ) if not value
        ]
        return {}, ["ไม่ได้กำหนด " + ", ".join(missing) + " ใน .env"]
    if max_duration_seconds is not None and max_duration_seconds <= 0:
        return {}, ["หมดเวลาประมวลผล Liquidity; ไม่ได้เริ่มการเชื่อมต่อ"]

    expected_errors = (OSError, RuntimeError, TimeoutError, ValueError)
    try:
        from requests.exceptions import RequestException
        from settrade_v2 import Investor
        from settrade_v2.errors import SettradeError
    except ImportError as exc:
        if investor_factory is None:
            return {}, [f"โหลด Settrade SDK ไม่สำเร็จ ({type(exc).__name__})"]
    else:
        expected_errors += (RequestException, SettradeError)
        if investor_factory is None:
            investor_factory = Investor

    broker_id = config.get("SETTRADE_BROKER_ID", "SANDBOX").strip() or "SANDBOX"
    app_code = config.get("SETTRADE_APP_CODE", "SANDBOX").strip() or "SANDBOX"
    results: dict[str, dict[str, Optional[float]]] = {}
    issues: list[str] = []
    started_at = time.monotonic()
    try:
        investor = investor_factory(
            app_id=app_id,
            app_secret=app_secret,
            app_code=app_code,
            broker_id=broker_id,
            is_auto_queue=False,
        )
    except expected_errors as exc:
        safe_message = str(exc).replace(app_id, "[REDACTED]").replace(app_secret, "[REDACTED]")
        return {}, [
            f"เชื่อมต่อ Settrade ไม่สำเร็จ ({type(exc).__name__}: {safe_message})"
        ]

    for symbol_index, symbol in enumerate(symbols):
        if max_duration_seconds is not None:
            remaining_duration = max_duration_seconds - (time.monotonic() - started_at)
            if remaining_duration <= 0:
                skipped = len(symbols) - symbol_index
                issues.append(
                    f"หมดเวลาประมวลผล Liquidity; ข้าม {skipped} หุ้นที่เหลือ"
                )
                break
        else:
            remaining_duration = SETTRADE_REALTIME_TIMEOUT_SECONDS

        state: dict[str, Any] = {"info": None, "book": None, "received": set()}
        state_lock = threading.Lock()
        received_data = threading.Event()
        callback_issues: list[str] = []
        subscribers = []

        def callback(kind: str, message: dict[str, Any]) -> None:
            try:
                parsed = _parse_market_message(message, kind)
            except (TypeError, ValueError) as exc:
                parsed = {}
                safe_message = str(exc).replace(app_id, "[REDACTED]").replace(app_secret, "[REDACTED]")
                callback_issues.append(f"{kind}: {safe_message}")
            with state_lock:
                state[kind] = parsed or None
                state["received"].add(kind)
                if len(state["received"]) == 2:
                    received_data.set()

        try:
            realtime = investor.RealtimeDataConnection()
            info_subscriber = realtime.subscribe_price_info(
                symbol, lambda message: callback("info", message)
            )
            subscribers.append(info_subscriber)
            info_subscriber.start()
            book_subscriber = realtime.subscribe_bid_offer(
                symbol, lambda message: callback("book", message)
            )
            subscribers.append(book_subscriber)
            book_subscriber.start()
            wait_seconds = SETTRADE_REALTIME_TIMEOUT_SECONDS
            if max_duration_seconds is not None:
                remaining_duration = max_duration_seconds - (
                    time.monotonic() - started_at
                )
                wait_seconds = min(wait_seconds, max(remaining_duration, 0))
            received_data.wait(wait_seconds)
        except expected_errors as exc:
            safe_message = str(exc).replace(app_id, "[REDACTED]").replace(app_secret, "[REDACTED]")
            issues.append(
                f"{symbol}: เชื่อมต่อ feed ไม่สำเร็จ ({type(exc).__name__}: {safe_message})"
            )
        finally:
            for subscriber in subscribers:
                try:
                    subscriber.stop()
                except (OSError, RuntimeError) as exc:
                    issues.append(
                        f"{symbol}: ปิด feed ไม่สำเร็จ ({type(exc).__name__})"
                    )

        with state_lock:
            info = state["info"]
            book = state["book"]
            if info or book:
                results[symbol] = {**(info or {}), **(book or {})}
            missing_feeds = [
                feed for feed in ("info", "book")
                if feed not in state["received"]
            ]
        issues.extend(f"{symbol}: {issue}" for issue in callback_issues)
        if missing_feeds:
            issues.append(
                f"{symbol}: ไม่ได้รับ feed " + ", ".join(missing_feeds)
            )
    return results, issues


def _collect_historical_turnover_estimates(
    symbols: list[str],
    config: dict[str, str],
    *,
    max_duration_seconds: Optional[float] = None,
    investor_factory: Optional[Callable[..., Any]] = None,
) -> tuple[dict[str, float], list[str]]:
    if not symbols:
        return {}, []
    app_id = config.get("SETTRADE_APP_ID", "").strip()
    app_secret = config.get("SETTRADE_APP_SECRET", "").strip()
    if not app_id or not app_secret:
        missing = [
            name for name, value in (
                ("SETTRADE_APP_ID", app_id),
                ("SETTRADE_APP_SECRET", app_secret),
            ) if not value
        ]
        return {}, ["Historical API: ไม่ได้กำหนด " + ", ".join(missing) + " ใน .env"]
    if max_duration_seconds is not None and max_duration_seconds <= 0:
        return {}, ["Historical API: หมดเวลาประมวลผลก่อนเริ่มดึงข้อมูล"]

    expected_errors = (OSError, RuntimeError, TimeoutError, ValueError)
    try:
        from requests.exceptions import RequestException
        from settrade_v2 import Investor
        from settrade_v2.errors import SettradeError
    except ImportError as exc:
        if investor_factory is None:
            return {}, [f"Historical API: โหลด Settrade SDK ไม่สำเร็จ ({type(exc).__name__})"]
    else:
        expected_errors += (RequestException, SettradeError)
        if investor_factory is None:
            investor_factory = Investor

    broker_id = config.get("SETTRADE_BROKER_ID", "SANDBOX").strip() or "SANDBOX"
    app_code = config.get("SETTRADE_APP_CODE", "SANDBOX").strip() or "SANDBOX"
    started_at = time.monotonic()
    try:
        investor = investor_factory(
            app_id=app_id,
            app_secret=app_secret,
            app_code=app_code,
            broker_id=broker_id,
            is_auto_queue=False,
        )
        market_data = investor.MarketData()
    except expected_errors as exc:
        safe_message = str(exc).replace(app_id, "[REDACTED]").replace(app_secret, "[REDACTED]")
        return {}, [
            f"Historical API: เชื่อมต่อ Settrade ไม่สำเร็จ ({type(exc).__name__}: {safe_message})"
        ]

    estimates: dict[str, float] = {}
    issues: list[str] = []
    for symbol_index, symbol in enumerate(symbols):
        if max_duration_seconds is not None:
            remaining = max_duration_seconds - (time.monotonic() - started_at)
            if remaining <= 0:
                issues.append(
                    f"Historical API: หมดเวลา; ข้าม {len(symbols) - symbol_index} หุ้นที่เหลือ"
                )
                break
        try:
            response = market_data.get_candlestick(
                symbol, "1d", limit=SETTRADE_HISTORY_CANDLE_LIMIT
            )
        except expected_errors as exc:
            safe_message = str(exc).replace(app_id, "[REDACTED]").replace(app_secret, "[REDACTED]")
            issues.append(
                f"{symbol}: Historical API ล้มเหลว ({type(exc).__name__}: {safe_message})"
            )
            continue

        fields = ("time", "close", "volume")
        if not isinstance(response, dict) or any(
            not isinstance(response.get(field), list) for field in fields
        ):
            issues.append(f"{symbol}: Historical API ส่งโครงสร้างแท่งรายวันไม่ครบ")
            continue
        lengths = {len(response[field]) for field in fields}
        if len(lengths) != 1:
            issues.append(f"{symbol}: จำนวนข้อมูล time/close/volume ไม่ตรงกัน")
            continue
        value_values = response.get("value")
        value_schema_invalid = not isinstance(value_values, list) or (
            len(value_values) != len(response["time"])
        )
        if value_schema_invalid:
            value_values = [None] * len(response["time"])

        today = datetime.now(THAILAND_TIMEZONE).date()
        sessions: list[tuple[float, float, float, Optional[float]]] = []
        invalid_timestamps = False
        for timestamp, close_value, volume_value, traded_value in zip(
            response["time"],
            response["close"],
            response["volume"],
            value_values,
        ):
            timestamp_number = _finite_number(timestamp)
            if timestamp_number is None or timestamp_number <= 0:
                invalid_timestamps = True
                continue
            try:
                session_date = datetime.fromtimestamp(
                    timestamp_number, timezone.utc
                ).astimezone(THAILAND_TIMEZONE).date()
            except (OverflowError, OSError, ValueError):
                invalid_timestamps = True
                continue
            volume = _finite_number(volume_value)
            if session_date >= today or volume is None or volume <= 0:
                continue
            close = _finite_number(close_value)
            if close is None or close <= 0:
                estimate = math.nan
            else:
                estimate = close * volume
            sessions.append(
                (timestamp_number, estimate, volume, _finite_number(traded_value))
            )

        if invalid_timestamps:
            issues.append(f"{symbol}: มี timestamp ใน Historical API ที่อ่านไม่ได้")
        sessions.sort(key=lambda session: session[0])
        sessions = sessions[-SETTRADE_HISTORY_TRADING_DAYS:]
        if len(sessions) < SETTRADE_HISTORY_TRADING_DAYS:
            issues.append(
                f"{symbol}: มีวันซื้อขายย้อนหลังที่ใช้ได้เพียง {len(sessions)}/"
                f"{SETTRADE_HISTORY_TRADING_DAYS} วัน"
            )
            continue
        if any(
            not math.isfinite(estimate) or estimate <= 0
            for _, estimate, _, _ in sessions
        ):
            issues.append(f"{symbol}: ข้อมูล close/volume ย้อนหลังไม่ครบสำหรับคำนวณ")
            continue

        estimates[symbol] = sum(
            estimate for _, estimate, _, _ in sessions
        ) / len(sessions)
        value_issues = sum(
            value is None or value <= 0 for _, _, _, value in sessions
        )
        if value_issues or value_schema_invalid:
            issues.append(
                f"{symbol}: ฟิลด์ value ย้อนหลังผิดปกติ "
                f"{value_issues}/{len(sessions)} วัน"
                + (" หรือโครงสร้างไม่ครบ" if value_schema_invalid else "")
                + "; "
                "ใช้ค่า close × volume เป็นค่าประมาณแทน"
            )
    return estimates, issues


def _classify_update_rows(rows: list, indexes: dict[str, int]) -> dict[str, list]:
    categories: dict[str, list] = {
        "BUY_CONFIRMED": [],
        "EARLY_WATCH": [],
        "PULLBACK_ZONE": [],
        "OVEREXTENDED": [],
        "WAIT_RETEST": [],
    }
    for row in rows:
        categories[_candidate_report_state(row, indexes)].append(row)

    macd_order = lambda row: -(_finite_number(_cell(row, indexes, "MACD_HIST_DELTA(W)")) or 0)
    pullback_order = lambda row: -_candidate_metrics(row, indexes)[3]
    for key in ("BUY_CONFIRMED", "EARLY_WATCH"):
        categories[key].sort(key=macd_order)
    categories["PULLBACK_ZONE"].sort(key=pullback_order)
    return categories


def _collect_update_liquidity(
    rows: list,
    indexes: dict[str, int],
    config: dict[str, str],
) -> tuple[dict[str, dict[str, Optional[float]]], list[tuple[str, int, int]], list[str]]:
    categories = _classify_update_rows(rows, indexes)
    displayed_categories = (
        ("BUY CONFIRMED", "BUY_CONFIRMED", None),
        ("EARLY/WATCH", "EARLY_WATCH", 10),
        ("PULLBACK ZONE", "PULLBACK_ZONE", 10),
        ("OVEREXTENDED", "OVEREXTENDED", None),
    )
    started_at = time.monotonic()
    all_liquidity: dict[str, dict[str, Optional[float]]] = {}
    history_symbols: list[str] = []
    round_summaries: list[tuple[str, int, int]] = []
    all_issues: list[str] = []
    missing_credentials = [
        key for key in ("SETTRADE_APP_ID", "SETTRADE_APP_SECRET")
        if not config.get(key, "").strip()
    ]
    for title, key, limit in displayed_categories:
        category_rows = categories[key]
        visible_rows = category_rows[:limit] if limit is not None else category_rows
        if not visible_rows:
            continue

        symbols = [
            _display(_cell(row, indexes, "AKA")).upper()
            for row in visible_rows
        ]
        history_symbols.extend(symbols)
        if missing_credentials:
            liquidity = {}
            issues = [
                "ไม่ได้กำหนด " + ", ".join(missing_credentials) + " ใน .env"
            ]
        else:
            remaining = max(
                SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS
                - (time.monotonic() - started_at),
                0,
            )
            liquidity, issues = _collect_settrade_liquidity(
                symbols, config, max_duration_seconds=remaining
            )
        all_liquidity.update(liquidity)
        required_metric_keys = {
            "turnover", "bid_price", "ask_price", "spread", "bid_depth", "ask_depth",
        }
        complete_count = sum(
            required_metric_keys.issubset(metrics)
            for metrics in liquidity.values()
        )
        round_summaries.append((title, complete_count, len(symbols)))
        all_issues.extend(f"{title}: {issue}" for issue in issues)

    history_symbols = list(dict.fromkeys(history_symbols))
    remaining = max(
        SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS - (time.monotonic() - started_at),
        0,
    )
    historical_estimates, historical_issues = _collect_historical_turnover_estimates(
        history_symbols, config, max_duration_seconds=remaining
    )
    for symbol, estimate in historical_estimates.items():
        all_liquidity.setdefault(symbol, {})["avg_turnover_20d_estimate"] = estimate
    all_issues.extend(
        f"Historical turnover warning: {issue}" for issue in historical_issues
    )
    return all_liquidity, round_summaries, all_issues


def _print_update_liquidity_summary(
    round_summaries: list[tuple[str, int, int]],
) -> None:
    for title, complete_count, total_count in round_summaries:
        print(f"Liquidity feed ครบ — {title}: {complete_count}/{total_count} หุ้น")


def _format_liquidity(metrics: Optional[dict[str, Optional[float]]]) -> str:
    if not metrics:
        return "⚫ N/A"
    turnover = metrics.get("turnover")
    bid_price = metrics.get("bid_price")
    ask_price = metrics.get("ask_price")
    spread = metrics.get("spread")
    bid_depth = metrics.get("bid_depth")
    ask_depth = metrics.get("ask_depth")
    return (
        f"Turnover ฿{turnover:,.2f}" if turnover is not None else "Turnover N/A"
    ) + " | " + (
        f"Bid1 ฿{bid_price:,.4f}" if bid_price is not None else "Bid1 N/A"
    ) + " / " + (
        f"Ask1 ฿{ask_price:,.4f}" if ask_price is not None else "Ask1 N/A"
    ) + " | " + (
        f"Spread ฿{spread:,.4f}" if spread is not None else "Spread N/A"
    ) + " | " + (
        f"Bid10 {bid_depth:,.0f}" if bid_depth is not None else "Bid10 N/A"
    ) + " / " + (
        f"Ask10 {ask_depth:,.0f}" if ask_depth is not None else "Ask10 N/A"
    )


def _format_historical_turnover_estimate(
    metrics: Optional[dict[str, Optional[float]]],
) -> str:
    estimate = _finite_number((metrics or {}).get("avg_turnover_20d_estimate"))
    return f"≈ ฿{estimate:,.2f}" if estimate is not None else "⚫ N/A"


def _print_update_report(
    worksheet,
    indexes: dict[str, int],
    selected_symbols: list[str],
    liquidity: Optional[dict[str, dict[str, Optional[float]]]] = None,
) -> bool:
    if not selected_symbols:
        rows = [
            row for row in worksheet.iter_rows(min_row=2, values_only=True)
            if row and any(value is not None and str(value).strip() for value in row)
        ]
        _print_update_status_report(rows, indexes, liquidity or {})
        return False

    rows_by_symbol = {
        _display(_cell(row, indexes, "AKA")).upper(): row
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row and any(value is not None and str(value).strip() for value in row)
    }
    specified_rows = [
        (symbol, rows_by_symbol[symbol])
        for symbol in selected_symbols
        if symbol in rows_by_symbol
    ]
    missing_symbols = [symbol for symbol in selected_symbols if symbol not in rows_by_symbol]
    incomplete_symbols = []

    for symbol, row in specified_rows:
        missing_values = [
            column for column in ("LAST_CANDLE", "UPDATED")
            if not _display(_cell(row, indexes, column))
        ]
        if missing_values:
            incomplete_symbols.append(symbol)
        print(f"AKA         : {symbol}")
        print(f"LAST_CANDLE : {_display(_cell(row, indexes, 'LAST_CANDLE'))}")
        print(f"UPDATED     : {_display(_cell(row, indexes, 'UPDATED'))}")
        print(f"GROUP       : {_display(_cell(row, indexes, 'GROUP'))}")
        print("")

    print(f"🔵 SPECIFIED STOCK — {len(specified_rows)} รายการ")
    columns = (
        "AKA", "GROUP", "RSI(W)", "RSI(1H)", "ราคา",
        "VOL (x/4)", "TREND (x/4)", "CMB (x/4)", "Liquidity",
        "Avg Turnover 20D*",
    )
    print("| " + " | ".join(columns) + " |")
    print("|---|---|---:|---:|---:|---|---|---|---|---:|")
    scored_rows = [
        (symbol, row, _candidate_metrics(row, indexes))
        for symbol, row in specified_rows
    ]
    scored_rows.sort(key=lambda item: -item[2][3])
    if not scored_rows:
        print("| " + " | ".join("—" for _ in columns) + " |")
    for symbol, row, metrics in scored_rows:
        volume, trend, combined, _ = metrics
        values = (
            f"[{symbol}]({finnomena_stock_url(symbol)})",
            _display(_cell(row, indexes, "GROUP")),
            _display(_cell(row, indexes, "RSI(W)")),
            _display(_cell(row, indexes, "RSI(1H)")),
            _format_price_fixed(_cell(row, indexes, "PRICE")),
            volume,
            trend,
            combined,
            _format_liquidity((liquidity or {}).get(symbol)),
            _format_historical_turnover_estimate((liquidity or {}).get(symbol)),
        )
        print("| " + " | ".join(
            value if index == 0 else _markdown_cell(value)
            for index, value in enumerate(values)
        ) + " |")

    if missing_symbols:
        print("\nไม่พบใน Snapshot ล่าสุด: " + ", ".join(missing_symbols))
    if incomplete_symbols:
        print(
            "\n⚠️ ไม่มีค่า LAST_CANDLE หรือ UPDATED ใน Snapshot ล่าสุด: "
            + ", ".join(incomplete_symbols)
        )
    return bool(missing_symbols or incomplete_symbols)


def _has_buy_volume(row, indexes: dict[str, int], timeframe: str) -> bool:
    return bool(re.search(
        r"\bBUY\b",
        _display(_cell(row, indexes, f"SIGNAL_VOLUME({timeframe})")),
        re.IGNORECASE,
    ))


def _candidate_metrics(row, indexes: dict[str, int]) -> tuple[str, str, str, int]:
    volume = [
        "🟢" if _has_buy_volume(row, indexes, timeframe) else "⚪"
        for timeframe in ("W", "D", "4H", "1H")
    ]
    trends = [
        "🟢" if _display(_cell(row, indexes, f"TREND({timeframe})")).upper() == "UP" else "🔴"
        for timeframe in ("W", "D", "4H", "1H")
    ]
    combined = [
        "✅" if volume_signal == "🟢" and trend_signal == "🟢" else "❌"
        for volume_signal, trend_signal in zip(volume, trends)
    ]
    score = combined.count("✅")
    return (
        f"{''.join(volume)} ({volume.count('🟢')}/4)",
        f"{''.join(trends)} ({trends.count('🟢')}/4)",
        f"{''.join(combined)} ({score}/4)",
        score,
    )


def _candidate_report_state(row, indexes: dict[str, int]) -> str:
    rsi_d = _finite_number(_cell(row, indexes, "RSI(D)"))
    trigger = _display(_cell(row, indexes, "ENTRY_TRIGGER(1H)")).upper()
    trigger_tokens = set(re.findall(r"[A-Z]+", trigger))
    trend_w = _display(_cell(row, indexes, "TREND(W)")).upper()
    trend_d = _display(_cell(row, indexes, "TREND(D)")).upper()
    trend_4h = _display(_cell(row, indexes, "TREND(4H)")).upper()
    trend_1h = _display(_cell(row, indexes, "TREND(1H)")).upper()
    rsi_1h = _finite_number(_cell(row, indexes, "RSI(1H)"))

    if rsi_d is not None and rsi_d >= 70:
        return "OVEREXTENDED"
    if (
        _candidate_state(row, indexes)
        and trigger_tokens.intersection({"BUY", "TRIGGER", "ENTRY"})
        and not trigger_tokens.intersection({"WAIT", "RETEST", "NO"})
    ):
        return "BUY_CONFIRMED"
    if (
        (trend_w, trend_d, trend_4h, trend_1h) == ("UP", "UP", "UP", "UP")
        and rsi_1h is not None
        and rsi_1h < 50
    ):
        return "PULLBACK_ZONE"
    if _candidate_state(row, indexes):
        return "EARLY_WATCH"
    return "WAIT_RETEST"


def _print_status_table(
    title: str,
    rows: list,
    indexes: dict[str, int],
    columns: tuple[str, ...],
    *,
    limit: Optional[int] = None,
    liquidity: Optional[dict[str, dict[str, Optional[float]]]] = None,
) -> None:
    print(f"\n{title}")
    print("| " + " | ".join(columns) + " |")
    print("|" + "|".join("---" for _ in columns) + "|")
    shown = rows[:limit] if limit is not None else rows
    if not shown:
        print("| " + " | ".join("—" for _ in columns) + " |")
        return

    for row in shown:
        symbol = _display(_cell(row, indexes, "AKA")).upper()
        values = []
        for column in columns:
            if column == "หุ้น":
                value = f"[{symbol}]({finnomena_stock_url(symbol)})"
            elif column == "GROUP":
                value = _display(_cell(row, indexes, "GROUP"))
            elif column == "RSI(W)":
                value = _display(_cell(row, indexes, "RSI(W)"))
            elif column == "RSI(D)":
                value = _display(_cell(row, indexes, "RSI(D)"))
            elif column == "RSI(1H)":
                value = _display(_cell(row, indexes, "RSI(1H)"))
            elif column == "MACD_D":
                value = _display(_cell(row, indexes, "MACD_HIST_DELTA(W)"))
            elif column == "ราคา":
                value = _format_price(_cell(row, indexes, "PRICE"))
            elif column == "VOL (x/4)":
                value = _candidate_metrics(row, indexes)[0]
            elif column == "TREND (x/4)":
                value = _candidate_metrics(row, indexes)[1]
            elif column == "CMB (x/4)":
                value = _candidate_metrics(row, indexes)[2]
            elif column == "Liquidity":
                value = _format_liquidity(
                    (liquidity or {}).get(symbol)
                )
            elif column == "Avg Turnover 20D*":
                value = _format_historical_turnover_estimate(
                    (liquidity or {}).get(symbol)
                )
            else:
                raise ValueError(f"ไม่รองรับคอลัมน์รายงาน: {column}")
            values.append(value if column == "หุ้น" else _markdown_cell(value))
        print("| " + " | ".join(values) + " |")


def _print_update_status_report(
    rows: list,
    indexes: dict[str, int],
    liquidity: Optional[dict[str, dict[str, Optional[float]]]] = None,
) -> None:
    missing = [column for column in REPORT_COLUMNS if column not in indexes]
    if missing:
        raise ValueError("ขาดคอลัมน์สำหรับรายงาน UPDATE: " + ", ".join(missing))

    categories = _classify_update_rows(rows, indexes)

    tabular_columns = (
        "หุ้น", "GROUP", "RSI(W)", "RSI(D)", "MACD_D", "ราคา", "Liquidity",
        "Avg Turnover 20D*",
    )
    _print_status_table(
        f"🟢 BUY CONFIRMED — {len(categories['BUY_CONFIRMED'])} หุ้น",
        categories["BUY_CONFIRMED"], indexes, tabular_columns, liquidity=liquidity,
    )
    _print_status_table(
        f"🟡 EARLY/WATCH — {len(categories['EARLY_WATCH'])} หุ้น (TOP 10)",
        categories["EARLY_WATCH"], indexes, tabular_columns, limit=10, liquidity=liquidity,
    )
    _print_status_table(
        f"🟠 PULLBACK ZONE — {len(categories['PULLBACK_ZONE'])} หุ้น (TOP 10)",
        categories["PULLBACK_ZONE"], indexes,
        ("หุ้น", "GROUP", "RSI(W)", "RSI(1H)", "ราคา", "VOL (x/4)", "TREND (x/4)", "CMB (x/4)", "Liquidity", "Avg Turnover 20D*"),
        limit=10, liquidity=liquidity,
    )
    _print_status_table(
        f"🔴 OVEREXTENDED — {len(categories['OVEREXTENDED'])} หุ้น",
        categories["OVEREXTENDED"], indexes,
        ("หุ้น", "GROUP", "Liquidity", "Avg Turnover 20D*"), liquidity=liquidity,
    )
    print(f"\n⚪ WAIT/RETEST — {len(categories['WAIT_RETEST'])} หุ้น")


def _candidate_buy_timeframes(row, indexes: dict[str, int]) -> list[str]:
    timeframes = ("W", "D", "4H", "1H")
    return [
        timeframe
        for timeframe in timeframes
        if re.search(r"\bBUY\b", _display(_cell(row, indexes, f"SIGNAL_VOLUME({timeframe})")), re.IGNORECASE)
    ]


def _candidate_state(row, indexes: dict[str, int]) -> bool:
    trend_w = _display(_cell(row, indexes, "TREND(W)")).upper()
    trend_d = _display(_cell(row, indexes, "TREND(D)")).upper()
    rsi_w = _finite_number(_cell(row, indexes, "RSI(W)"))
    macd_delta_w = _finite_number(_cell(row, indexes, "MACD_HIST_DELTA(W)"))
    rsi_d = _finite_number(_cell(row, indexes, "RSI(D)"))
    ema9_w = _finite_number(_cell(row, indexes, "EMA9(W)"))
    ema21_w = _finite_number(_cell(row, indexes, "EMA21(W)"))

    return (
        trend_w == "UP"
        and trend_d == "UP"
        and rsi_w is not None and rsi_w > 50
        and macd_delta_w is not None and macd_delta_w > 0
        and rsi_d is not None and rsi_d < 70
        and ema9_w is not None and ema21_w is not None and ema9_w > ema21_w
        and bool(_candidate_buy_timeframes(row, indexes))
    )


def _print_candidate_report(
    worksheet,
    indexes: dict[str, int],
    liquidity: Optional[dict[str, dict[str, Optional[float]]]] = None,
) -> None:
    missing = [column for column in REPORT_COLUMNS if column not in indexes]
    if missing:
        raise ValueError(
            "ขาดคอลัมน์สำหรับ CANDIDATE: " + ", ".join(missing)
        )

    candidates = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue
        for column in ("EMA9(W)", "EMA21(W)"):
            if _finite_number(_cell(row, indexes, column)) is None:
                raise ValueError(f"แถว {row_number}: {column} ต้องเป็นตัวเลข finite")
        if _candidate_state(row, indexes):
            candidates.append(row)

    candidates.sort(
        key=lambda row: -(_finite_number(_cell(row, indexes, "MACD_HIST_DELTA(W)")) or 0)
    )
    print(f"\nผลการค้นหา CANDIDATE (PRE-ENTRY-Like Pattern) — {len(candidates)} หุ้น")
    print("ผ่านเกณฑ์ CANDIDATE 7 ข้อ เรียงตาม MACD Delta (W) สูงสุด — รอ 1H Trigger ยืนยัน")
    columns = (
        "#", "หุ้น", "GROUP", "ราคา (THB)", "MACD Δ(W)",
        "Entry Zone (4H)", "1H Trigger", "Liquidity", "Avg Turnover 20D*",
    )
    print("| " + " | ".join(columns) + " |")
    print("|---:|---|---|---:|---:|---|---|---|---:|")
    if not candidates:
        print("| " + " | ".join("—" for _ in columns) + " |")
        return
    for number, row in enumerate(candidates[:20], start=1):
        symbol = _display(_cell(row, indexes, "AKA")).upper()
        values = [
            str(number),
            f"[{symbol}]({finnomena_stock_url(symbol)})",
            _display(_cell(row, indexes, "GROUP")),
            _format_price(_cell(row, indexes, "PRICE")),
            _display(_cell(row, indexes, "MACD_HIST_DELTA(W)")),
            _display(_cell(row, indexes, "ENTRY_ZONE(4H)")),
            _display(_cell(row, indexes, "ENTRY_TRIGGER(1H)")),
            _format_liquidity((liquidity or {}).get(symbol)),
            _format_historical_turnover_estimate((liquidity or {}).get(symbol)),
        ]
        print("| " + " | ".join(
            value if index == 1 else _markdown_cell(value)
            for index, value in enumerate(values)
        ) + " |")


def main(argv: Optional[list[str]] = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    command = args.pop(0).upper() if args else "UPDATE"
    if command not in {"UPDATE", "CANDIDATE"}:
        if command.startswith("UPDATE/"):
            args.insert(0, command[len("UPDATE"):])
            command = "UPDATE"
        else:
            print("Usage: set_update.py [UPDATE[/AKA-AKA-...] | CANDIDATE]", file=sys.stderr)
            return 2
    selected_symbols: list[str] = []
    if command == "CANDIDATE" and args:
        print("CANDIDATE ไม่รับรายการ AKA เพิ่ม", file=sys.stderr)
        return 2
    if command == "UPDATE":
        if len(args) > 1:
            print("Usage: set_update.py [UPDATE[/AKA-AKA-...]]", file=sys.stderr)
            return 2
        if args:
            raw = args[0]
            if not raw.startswith("/") or not raw[1:]:
                print("รูปแบบ UPDATE ที่รองรับคือ UPDATE/<AKA>-<AKA>-...", file=sys.stderr)
                return 2
            selected_symbols = [symbol.upper() for symbol in raw[1:].split("-")]
            if any(not SYMBOL_PATTERN.fullmatch(symbol) for symbol in selected_symbols):
                print("รายการ AKA ใน UPDATE มีรูปแบบไม่ถูกต้อง", file=sys.stderr)
                return 2
            selected_symbols = list(dict.fromkeys(selected_symbols))

    try:
        config = load_config()
        export_url = spreadsheet_export_url(config["SET_DATA_URL"])
        download_workbook(export_url)
        workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    except (ValueError, BadZipFile, OSError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        print(f"❌ {exc}", file=sys.stderr)
        return 1

    sheet_name = config["SET_DATA_SHEET"]
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        print(f"❌ ไม่พบชีต {sheet_name!r}; ชีตที่มี: {available}", file=sys.stderr)
        return 1

    try:
        worksheet = workbook[sheet_name]
        row_count, errors, timestamp_available, error_count = validate_worksheet(worksheet)
        print(f"ชีต: {sheet_name}")
        print(f"จำนวนแถวข้อมูล: {row_count}")

        if errors:
            print(f"❌ พบข้อผิดพลาด {error_count} รายการ:")
            for error in errors[:20]:
                print(f"  - {error}")
            if error_count > 20:
                print(f"  - และข้อผิดพลาดอื่นอีก {error_count - 20} รายการ")
            return 1

        indexes = _header_indexes(worksheet)
        missing = [column for column in REPORT_COLUMNS if column not in indexes]
        if missing:
            command_label = "CANDIDATE" if command == "CANDIDATE" else "UPDATE"
            print(
                f"❌ ขาดคอลัมน์สำหรับรายงาน {command_label}: " + ", ".join(missing),
                file=sys.stderr,
            )
            return 1
        if command == "UPDATE" and not selected_symbols:
            summary_columns = ("LAST_CANDLE", "UPDATED")
            missing_summary_columns = [
                column for column in summary_columns if column not in indexes
            ]
            if missing_summary_columns:
                print(
                    "❌ ขาดคอลัมน์สำหรับสรุป UPDATE: "
                    + ", ".join(missing_summary_columns),
                    file=sys.stderr,
                )
                return 1
        if command == "UPDATE" and selected_symbols:
            missing_summary_columns = [
                column for column in ("LAST_CANDLE", "UPDATED")
                if column not in indexes
            ]
            if missing_summary_columns:
                print(
                    "❌ ขาดคอลัมน์สำหรับ UPDATE รายการที่ระบุ: "
                    + ", ".join(missing_summary_columns),
                    file=sys.stderr,
                )
                return 1

        print(f"✅ ดาวน์โหลด snapshot: {WORKBOOK_PATH}")
        print("✅ โครงสร้างและชนิดข้อมูลพื้นฐานผ่านการตรวจสอบ")
        update_summary_row = None
        if command == "UPDATE" and not selected_symbols:
            try:
                update_summary_row = _update_summary_row(worksheet, indexes)
            except ValueError as exc:
                print(f"❌ {exc}", file=sys.stderr)
                return 1
        rows = [
            row for row in worksheet.iter_rows(min_row=2, values_only=True)
            if row and any(value is not None and str(value).strip() for value in row)
        ]
        update_liquidity_summaries = []
        if command == "UPDATE" and selected_symbols:
            selected_set = set(selected_symbols)
            symbols = [
                _display(_cell(row, indexes, "AKA")).upper()
                for row in rows
                if _display(_cell(row, indexes, "AKA")).upper() in selected_set
            ]
            liquidity_started_at = time.monotonic()
            liquidity, liquidity_issues = _collect_settrade_liquidity(
                list(dict.fromkeys(symbols)),
                config,
                max_duration_seconds=SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS,
            )
            remaining = max(
                SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS
                - (time.monotonic() - liquidity_started_at),
                0,
            )
            estimates, historical_issues = _collect_historical_turnover_estimates(
                list(dict.fromkeys(symbols)), config, max_duration_seconds=remaining
            )
            for symbol, estimate in estimates.items():
                liquidity.setdefault(symbol, {})[
                    "avg_turnover_20d_estimate"
                ] = estimate
            liquidity_issues.extend(
                f"Historical turnover warning: {issue}"
                for issue in historical_issues
            )
        elif command == "CANDIDATE":
            candidate_rows = [
                row for row in rows if _candidate_state(row, indexes)
            ]
            candidate_rows.sort(
                key=lambda row: -(
                    _finite_number(_cell(row, indexes, "MACD_HIST_DELTA(W)")) or 0
                )
            )
            candidate_symbols = [
                _display(_cell(row, indexes, "AKA")).upper()
                for row in candidate_rows[:20]
            ]
            liquidity_started_at = time.monotonic()
            liquidity, liquidity_issues = _collect_settrade_liquidity(
                candidate_symbols,
                config,
                max_duration_seconds=SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS,
            )
            remaining = max(
                SETTRADE_LIQUIDITY_MAX_DURATION_SECONDS
                - (time.monotonic() - liquidity_started_at),
                0,
            )
            estimates, historical_issues = _collect_historical_turnover_estimates(
                candidate_symbols, config, max_duration_seconds=remaining
            )
            for symbol, estimate in estimates.items():
                liquidity.setdefault(symbol, {})[
                    "avg_turnover_20d_estimate"
                ] = estimate
            liquidity_issues.extend(
                f"Historical turnover warning: {issue}"
                for issue in historical_issues
            )
        elif command == "UPDATE":
            (
                liquidity,
                update_liquidity_summaries,
                liquidity_issues,
            ) = _collect_update_liquidity(rows, indexes, config)
        else:
            liquidity, liquidity_issues = {}, []
            update_liquidity_summaries = []
        try:
            if command == "UPDATE":
                if not selected_symbols:
                    _print_update_summary(update_summary_row, indexes)
                missing_symbols = _print_update_report(
                    worksheet, indexes, selected_symbols, liquidity
                )
            else:
                _print_candidate_report(worksheet, indexes, liquidity)
        except ValueError as exc:
            print(f"❌ {exc}", file=sys.stderr)
            return 1
        if update_liquidity_summaries:
            _print_update_liquidity_summary(update_liquidity_summaries)
        if command in {"UPDATE", "CANDIDATE"}:
            print(
                "⚠️ Avg Turnover 20D* เป็นค่าประมาณจากราคาปิด × volume "
                "ของ 20 วันซื้อขายที่มี volume ก่อนวันนี้ ไม่ใช่มูลค่าซื้อขายจริง"
            )
        if not timestamp_available:
            print("⚠️ ไม่พบคอลัมน์เวลาอ้างอิง จึงตรวจความสดของข้อมูลไม่ได้")
        else:
            print("⚠️ พบคอลัมน์เวลาอ้างอิง แต่ยังไม่ได้ตรวจความสดของข้อมูล")
        for issue in liquidity_issues:
            if issue.startswith("Historical turnover warning:"):
                print(f"⚠️ Historical data: {issue.removeprefix('Historical turnover warning: ')}")
            else:
                print(f"⚠️ Liquidity unavailable: {issue}")
        print("⚠️ ไม่ได้ตรวจสอบ benchmark, ปฏิทิน/เวลาซื้อขาย, corporate actions, tick size หรือ board lot")
        print("⚠️ ใช้ signals ที่มีใน snapshot เป็นข้อมูลวิเคราะห์; ไม่คำนวณ OHLCV ใหม่และไม่ส่งคำสั่งซื้อขาย")
        if command == "UPDATE" and missing_symbols:
            return 1
        return 0
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
